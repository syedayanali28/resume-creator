"""
Generate application materials under:

  people/<Person>/<Company>/<Role>/
    <Person>_<Company>_<Role>_resume.tex
    <Person>_<Company>_<Role>_resume.pdf
    <Person>_<Company>_<Role>_cover-letter.tex
    <Person>_<Company>_<Role>_cover-letter.pdf
    job link.txt          (posting URL; create/update with --job-url)

Optional Word exports with --also-docx.

Master / general-purpose resume lives at:

  people/<Person>/complete/
    master-resume.tex
    resume-knowledge-base.yaml
    master-resume.pdf   (run compile_complete_master.py)

Examples:
  python generate_job_package.py --person Ali-Syed-Ayan --company Creaxon-Technologies --role Software-Engineer-HK --job-url "https://..."
  python compile_complete_master.py --person Ali-Syed-Ayan
  python compile_complete_master.py --person Zhao-Yanbo
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

_SCRIPT_DIR = Path(__file__).resolve().parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from layout_paths import (
    ensure_layout,
    job_packet_dir,
    tailored_cover_letter_tex,
    tailored_job_link_txt,
    tailored_packet_basename,
    tailored_resume_tex,
)
from latex_build import pdflatex

# Optional Word pipeline (same tailored content as before)
try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    from resume_docx_template import (
        BODY_PT,
        FONT_NAME,
        NAME_PT,
        add_bullets,
        add_job_meta_line,
        add_job_spacer,
        add_job_title_line,
        add_name_header,
        add_section_heading,
        apply_body_styles,
        apply_resume_page_setup,
        endash_range,
        polish_all_runs_times,
        scrub,
        _set_run_font,
    )

    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

_BANNED = None
if _HAS_DOCX:
    import re

    _BANNED = re.compile(r"[;\u2014]")


def _validate_resume_text(lines: list[str]) -> None:
    if not _HAS_DOCX or _BANNED is None:
        return
    for line in lines:
        cleaned = scrub(line)
        if _BANNED.search(cleaned):
            raise ValueError(f"Resume text still contains banned punctuation: {cleaned!r}")


def docx_to_pdf_word(docx_path: Path, pdf_path: Path) -> None:
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        try:
            doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        finally:
            doc.Close(SaveChanges=False)
    finally:
        word.Quit()


def build_resume_qrt() -> "Document":
    doc = Document()
    apply_resume_page_setup(doc)
    apply_body_styles(doc)

    contact = (
        "+852 97537220  \u00b7  syedayanali28@gmail.com  \u00b7  "
        "https://github.com/syedayanali28  \u00b7  https://www.linkedin.com/in/syedayanali/"
    )
    add_name_header(doc, "Ali Syed Ayan", contact)

    add_section_heading(doc, "Education", first=True)
    add_job_title_line(doc, "City University of Hong Kong", endash_range("Jun 2022", "Jun 2026"))
    add_job_meta_line(
        doc,
        "Bachelor of Science in Computer Science, Minor in Business Intelligence",
        "Hong Kong",
    )
    edu_bullets = [
        "Scholarships and awards cover International Student Top Scholarship HK$720,000 total plus Dean's List and Tiger Talent Program 3 times.",
        "Coursework spans Web Application Development, Machine Learning, Algorithms, Data Structures, Database Systems, Computer Vision, Data Science, Networks, and OS.",
    ]
    _validate_resume_text(edu_bullets)
    add_bullets(doc, edu_bullets)

    add_section_heading(doc, "Work Experience")

    add_job_title_line(doc, "Hong Kong Monetary Authority", endash_range("Jul 2025", "Present"))
    add_job_meta_line(
        doc,
        "Analyst Programmer (IT Division \u2013 Business Enablement and Software Architecture Team)",
        "Hong Kong",
    )
    hkma = [
        "BRDR teams faced blind spots on about 10,000 PDFs so I built Python RAG with pgvector, SQL vector search, and Grafana and cut manual doc hunts for daily ops.",
        "Firewall reviews took about 30 staff hours monthly so I shipped Next.js automation with JIRA and OPA and returned about 30 hours monthly to security.",
        "Leaders needed faster KPI views so I built Tableau and Power BI dashboards with PO-led specs and training that cut weekly reporting prep versus manual packs.",
    ]
    _validate_resume_text(hkma)
    add_bullets(doc, hkma)

    add_job_spacer(doc)
    add_job_title_line(doc, "Jay Cee Ess Global", endash_range("Aug 2025", "Present"))
    add_job_meta_line(doc, "Software Engineer", "Remote (Australia-based)")
    jce = [
        "Revenue flow stalled on a legacy UI so I rebuilt the web app and funnel and lifted user intake about 200% with gains tied to about AUD 100,000 monthly revenue context.",
        "Checkout drop off hurt sales so I added Stripe with Apple Pay, Google Pay, and Alipay and cut cart abandonment about 70%.",
        "Pricing calls ran slow so I coded in-memory Yen k-shortest-path routing with Redis and cut latency about 99% for quote pages.",
        "Supabase spend blocked scale so I moved hot paths to Upstash Redis, cut database cost about 75%, and added multi-carrier routes.",
    ]
    _validate_resume_text(jce)
    add_bullets(doc, jce)

    add_job_spacer(doc)
    add_job_title_line(doc, "Universal Innotech Group Company Limited", endash_range("Sep 2024", "May 2025"))
    add_job_meta_line(doc, "Junior Software Engineer", "Tsuen Wan, Hong Kong")
    uni = [
        "Open Sanction API bills hit about HK$10,000 monthly so I shipped Neo4j graph APIs and Python scrapers for AML and PEP and removed that recurring fee.",
        "Broker UI felt slow so I built React SSR, caching, and reusable components and cut load time over 50% on admin and ticket views.",
        "Clients needed faster AML insight so REST and Neo4j views cut client decision time about 30% versus prior static screens.",
    ]
    _validate_resume_text(uni)
    add_bullets(doc, uni)

    add_job_spacer(doc)
    add_job_title_line(doc, "Yi Shing Group Ltd. (YSG)", endash_range("Jun 2024", "Aug 2024"))
    add_job_meta_line(doc, "Machine Learning Intern", "Science Park, Hong Kong")
    ysg = [
        "Field tests needed a live web control plane so I shipped Flask to Vertex AI with CI/CD and held about 90% to 95% accuracy in deployment checks.",
    ]
    _validate_resume_text(ysg)
    add_bullets(doc, ysg)

    add_section_heading(doc, "Projects")
    proj = [
        "Users needed a self-custodial wallet so I built a Next.js TypeScript Web3 app with create, connect, and send flows.",
        "Movie search needed snappy filters so I shipped Flask REST APIs with pagination and genre, cast, and year filters for large catalogs.",
    ]
    _validate_resume_text(proj)
    add_bullets(doc, proj)

    add_section_heading(doc, "Technical Skills")
    skills_para = (
        "TypeScript, JavaScript, React, Next.js, HTML, CSS, Python, Flask, REST, Redis, SQL, "
        "Neo4j, Git, Grafana, Tableau, Power BI, AWS, GCP, Vercel, agile scrum, UI for data-heavy screens, API ties to back ends."
    )
    _validate_resume_text([skills_para])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(scrub(skills_para))
    _set_run_font(r, FONT_NAME, BODY_PT)

    polish_all_runs_times(doc)

    if doc.paragraphs:
        for r in doc.paragraphs[0].runs:
            r.font.size = Pt(NAME_PT)
            r.bold = True
    return doc


def build_cover_qrt() -> "Document":
    doc = Document()
    apply_resume_page_setup(doc)
    apply_body_styles(doc)

    p = doc.add_paragraph()
    r = p.add_run(
        scrub(
            "Ali Syed Ayan\n+852 97537220  \u00b7  syedayanali28@gmail.com  \u00b7  "
            "https://github.com/syedayanali28  \u00b7  https://www.linkedin.com/in/syedayanali/\n"
        )
    )
    _set_run_font(r, FONT_NAME, BODY_PT)

    p = doc.add_paragraph()
    r = p.add_run(scrub("12 April 2026\n\nHiring Team\nQube Research & Technologies\nHong Kong\n"))
    _set_run_font(r, FONT_NAME, BODY_PT)

    p = doc.add_paragraph()
    r = p.add_run(scrub("Re: Web GUI Developer - Quantitative Technology (Hong Kong)"))
    r.bold = True
    _set_run_font(r, FONT_NAME, BODY_PT)

    body = (
        "Dear Hiring Manager,\n\n"
        "I am writing to apply for the Web GUI Developer role in Hong Kong. QRT blends data, research, "
        "technology, and trading in one culture, and I want to build monitoring UIs that help teams decide fast "
        "under load.\n\n"
        "At the Hong Kong Monetary Authority I ship Next.js tooling and Grafana views on large retrieval stacks. "
        "I join daily scrum work with product owners, turn requests into concrete UI and reporting flows, and own "
        "production fixes when monitors fire.\n\n"
        "At Jay Cee Ess Global I rebuilt a legacy web app, cut pricing latency about 99% with in-memory routing "
        "and Redis, and cut cart loss about 70% after Stripe and wallet checkouts. At Universal Innotech I cut "
        "page load over 50% with React SSR and caching and cut API spend about HK$10,000 monthly with Neo4j graph "
        "work tied to AML screens.\n\n"
        "I am strongest in TypeScript, JavaScript, React, and Next.js, and I like tight loops with traders and "
        "researchers. I want to grow WebSockets and D3-style views in production with your stack.\n\n"
        "Thank you for your time. I would gladly discuss how I can help QRT ship fast, clear monitoring GUIs.\n\n"
        "Sincerely,\n\n"
        "Ali Syed Ayan"
    )
    p = doc.add_paragraph()
    r = p.add_run(scrub(body))
    _set_run_font(r, FONT_NAME, BODY_PT)

    polish_all_runs_times(doc)
    return doc


COMPANY_HANDLERS_DOCX = {
    "Qube-Research-and-Technologies": (build_resume_qrt, build_cover_qrt),
}


def main() -> int:
    parser = argparse.ArgumentParser(description="Build tailored resume + cover letter (LaTeX + PDF).")
    parser.add_argument("--person", default="Ali-Syed-Ayan")
    parser.add_argument(
        "--company",
        required=True,
        help="Company folder slug under people/<Person>/ (e.g. Creaxon-Technologies, Binance, HSBC).",
    )
    parser.add_argument(
        "--role",
        required=True,
        help="Role folder slug under people/<Person>/<Company>/ (e.g. Software-Engineer-HK).",
    )
    parser.add_argument(
        "--also-docx",
        action="store_true",
        help="Also write <Person>_<Company>_<Role>_resume.docx / _cover-letter.docx and PDF via Microsoft Word.",
    )
    parser.add_argument(
        "--job-url",
        default=None,
        metavar="URL",
        help="Write or overwrite job link.txt in the job folder with this posting URL (one line).",
    )
    args = parser.parse_args()

    ensure_layout(args.person, args.company, args.role)
    out_dir = job_packet_dir(args.person, args.company, args.role)
    bas = tailored_packet_basename(args.person, out_dir)

    job_link = tailored_job_link_txt(out_dir)
    if args.job_url is not None and str(args.job_url).strip():
        job_link.write_text(str(args.job_url).strip() + "\n", encoding="utf-8")

    resume_tex = tailored_resume_tex(args.person, out_dir)
    cover_tex = tailored_cover_letter_tex(args.person, out_dir)

    if not resume_tex.is_file() or not cover_tex.is_file():
        print(
            f"Missing {resume_tex.name} or {cover_tex.name} "
            f"(expected '{bas}_resume.tex' and '{bas}_cover-letter.tex'). Add LaTeX sources under:\n  {out_dir}",
            file=sys.stderr,
        )
        return 1

    try:
        pdflatex(resume_tex)
        pdflatex(cover_tex)
    except subprocess.CalledProcessError as e:
        print(e.stderr or e.stdout or str(e), file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(str(e), file=sys.stderr)
        return 1

    lines = [
        f"LaTeX + PDF:",
        f"  {resume_tex}",
        f"  {resume_tex.with_suffix('.pdf')}",
        f"  {cover_tex}",
        f"  {cover_tex.with_suffix('.pdf')}",
    ]
    if job_link.is_file():
        lines.extend(["Posting link file:", f"  {job_link}"])

    if args.also_docx:
        if not _HAS_DOCX:
            print("--also-docx requires python-docx.", file=sys.stderr)
            return 1
        pair = COMPANY_HANDLERS_DOCX.get(args.company)
        if not pair:
            print(f"No DOCX builder registered for company {args.company!r}.", file=sys.stderr)
            return 1
        br, bc = pair
        rp = out_dir / f"{bas}_resume.docx"
        cp = out_dir / f"{bas}_cover-letter.docx"
        br().save(rp)
        bc().save(cp)
        try:
            docx_to_pdf_word(rp, out_dir / f"{bas}_resume-word.pdf")
            docx_to_pdf_word(cp, out_dir / f"{bas}_cover-letter-word.pdf")
            lines.extend(
                [
                    f"  {rp}",
                    f"  {cp}",
                    f"  {out_dir / f'{bas}_resume-word.pdf'}",
                    f"  {out_dir / f'{bas}_cover-letter-word.pdf'}",
                ]
            )
        except Exception as e:
            print(f"Word PDF export failed ({e}). DOCX saved.", file=sys.stderr)

    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
